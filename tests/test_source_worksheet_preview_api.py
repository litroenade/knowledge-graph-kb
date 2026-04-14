from __future__ import annotations

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook


def _build_workbook_bytes() -> bytes:
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    generals = workbook.create_sheet("Generals")
    generals.append(["ID", "Name", "Type"])
    generals.append(["1", "Bai Qi", "General"])
    generals.append(["2", "Han Xin", "General"])
    generals.append(["3", "Sun Bin", "Strategist"])
    generals.append(["4", "Wu Qi", "General"])
    generals.append(["5", "Li Mu", "General"])

    terrain = workbook.create_sheet("Terrain")
    terrain.append(["Code", "Label"])
    terrain.append(["T1", "Mountain"])
    terrain.append(["T2", "River"])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _import_excel_source(client: TestClient) -> tuple[str, str]:
    response = client.post(
        "/api/kb/imports/uploads",
        files=[
            (
                "files",
                (
                    "military.xlsx",
                    _build_workbook_bytes(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            )
        ],
        data={"strategy": "summary"},
    )
    assert response.status_code == 200, response.text
    job = response.json()["job"]
    assert job["status"] in {"completed", "partial"}

    sources_response = client.get("/api/kb/sources")
    assert sources_response.status_code == 200, sources_response.text
    sources = sources_response.json()
    assert len(sources) == 1
    source = sources[0]
    assert source["active_version_id"]
    return source["id"], source["active_version_id"]


def test_source_worksheet_list_returns_sheet_level_contract(client: TestClient) -> None:
    source_id, version_id = _import_excel_source(client)

    response = client.get(
        f"/api/kb/sources/{source_id}/worksheets",
        params={"version_id": version_id},
    )
    assert response.status_code == 200, response.text

    payload = response.json()
    worksheet_by_key = {item["worksheet_key"]: item for item in payload["items"]}
    assert set(worksheet_by_key) == {"generals", "terrain"}

    generals = worksheet_by_key["generals"]
    assert generals == {
        "worksheet_key": "generals",
        "worksheet_name": "Generals",
        "row_count": 5,
        "headers": ["ID", "Name", "Type"],
        "column_keys": ["id", "name", "type"],
    }


def test_source_worksheet_preview_returns_pagination_contract(client: TestClient) -> None:
    source_id, version_id = _import_excel_source(client)

    response = client.get(
        f"/api/kb/sources/{source_id}/worksheets/generals/preview",
        params={"version_id": version_id, "page": 1, "page_size": 2},
    )
    assert response.status_code == 200, response.text

    payload = response.json()
    assert payload["worksheet_key"] == "generals"
    assert payload["worksheet_name"] == "Generals"
    assert payload["render_kind"] == "worksheet_preview"
    assert payload["headers"] == ["ID", "Name", "Type"]
    assert payload["column_keys"] == ["id", "name", "type"]
    assert payload["page"] == 1
    assert payload["page_size"] == 2
    assert payload["total_rows"] == 5
    assert payload["has_prev"] is False
    assert payload["has_next"] is True
    assert payload["row_range_start"] == 2
    assert payload["row_range_end"] == 3
    assert [item["row_index"] for item in payload["items"]] == [2, 3]
    assert payload["items"][0]["cells"] == {"id": "1", "name": "Bai Qi", "type": "General"}
    assert payload["render_metadata"]["page"] == 1
    assert payload["render_metadata"]["page_size"] == 2
    assert payload["render_metadata"]["total_rows"] == 5
    assert payload["render_metadata"]["has_prev"] is False
    assert payload["render_metadata"]["has_next"] is True
    assert "Generals" in (payload["rendered_html"] or "")


def test_source_worksheet_preview_anchor_row_resolves_context_window(client: TestClient) -> None:
    source_id, version_id = _import_excel_source(client)

    response = client.get(
        f"/api/kb/sources/{source_id}/worksheets/generals/preview",
        params=[
            ("version_id", version_id),
            ("page_size", "2"),
            ("anchor_row", "5"),
            ("highlighted_columns", "name"),
            ("highlighted_columns", "type"),
        ],
    )
    assert response.status_code == 200, response.text

    payload = response.json()
    assert payload["page"] == 2
    assert payload["page_size"] == 2
    assert payload["row_range_start"] == 4
    assert payload["row_range_end"] == 5
    assert [item["row_index"] for item in payload["items"]] == [4, 5]
    assert payload["anchor_row_index"] == 5
    assert payload["highlighted_row_indexes"] == [5]
    assert payload["highlighted_columns"] == ["name", "type"]
    assert payload["render_metadata"]["anchor_row_index"] == 5
    assert payload["render_metadata"]["highlighted_row_indexes"] == [5]
    assert payload["render_metadata"]["highlighted_columns"] == ["name", "type"]
    assert payload["render_metadata"]["row_range_start"] == 4
    assert payload["render_metadata"]["row_range_end"] == 5


def test_source_worksheet_preview_clamps_page_out_of_range(client: TestClient) -> None:
    source_id, version_id = _import_excel_source(client)

    response = client.get(
        f"/api/kb/sources/{source_id}/worksheets/generals/preview",
        params={"version_id": version_id, "page": 99, "page_size": 2},
    )
    assert response.status_code == 200, response.text

    payload = response.json()
    assert payload["page"] == 3
    assert payload["has_prev"] is True
    assert payload["has_next"] is False
    assert payload["row_range_start"] == 6
    assert payload["row_range_end"] == 6
    assert [item["row_index"] for item in payload["items"]] == [6]
    assert payload["render_metadata"]["fallback_reason"] == "page_out_of_range_clamped"


def test_source_worksheet_preview_rejects_unknown_worksheet(client: TestClient) -> None:
    source_id, version_id = _import_excel_source(client)

    response = client.get(
        f"/api/kb/sources/{source_id}/worksheets/unknown-sheet/preview",
        params={"version_id": version_id},
    )
    assert response.status_code == 404, response.text
    assert response.json() == {
        "code": "worksheet_not_found",
        "message": "Worksheet not found.",
    }
